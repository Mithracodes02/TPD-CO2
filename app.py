import io
import re
import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt
from matplotlib.ticker import AutoMinorLocator
from scipy.integrate import simpson

# Export Libraries
from docx import Document
from docx.shared import Inches, Pt
import reportlab
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

st.set_page_config(
    page_title="CO2-TPD Advanced Analyzer & Reporter",
    page_icon="🧪",
    layout="wide",
)

st.title("🧪 CO₂-TPD Advanced Analyzer, Benchmarker & Report Generator")
st.markdown(
    "Parse multi-block TPD files, customize publication graphics with rich typography, quantify basicity distribution, "
    "and export comprehensive reports in **PDF**, **Excel**, or **Word** formats."
)

def clean_numeric_series(series):
    """Converts a pandas series to numeric floats, handling text, units, and European commas."""
    def extract_num(val):
        if pd.isna(val):
            return np.nan
        val_str = str(val).strip().replace(',', '.')
        match = re.search(r'[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?', val_str)
        if match:
            try:
                return float(match.group(0))
            except ValueError:
                return np.nan
        return np.nan

    return series.apply(extract_num)


# --- EXPORT REPORT GENERATORS ---
def generate_excel_report(summary_df, metrics_dict, fig_buffer):
    """Creates a formatted Excel report with metrics, data table, and embedded plot."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Sheet 1: Summary & Metrics
        metrics_df = pd.DataFrame(list(metrics_dict.items()), columns=["Metric", "Value"])
        metrics_df.to_excel(writer, sheet_name="TPD Summary", index=False, startrow=0)
        summary_df.to_excel(writer, sheet_name="TPD Summary", index=False, startrow=len(metrics_df) + 3)
        
        # Add graph image to openpyxl sheet
        ws = writer.sheets["TPD Summary"]
        from openpyxl.drawing.image import Image as OPImage
        fig_buffer.seek(0)
        img = OPImage(fig_buffer)
        ws.add_image(img, "F2")
        
    return output.getvalue()


def generate_word_report(summary_df, metrics_dict, fig_buffer, conclusions_text):
    """Creates a Word document report containing graphics, tables, and literature insights."""
    doc = Document()
    doc.add_heading("CO₂-TPD Characterization Report", level=0)

    doc.add_heading("1. Experimental Metrics", level=1)
    for k, v in metrics_dict.items():
        doc.add_paragraph(f"• {k}: {v}")

    doc.add_heading("2. Basic Site Quantification", level=1)
    t = doc.add_table(rows=1, cols=len(summary_df.columns))
    hdr_cells = t.rows[0].cells
    for i, col_name in enumerate(summary_df.columns):
        hdr_cells[i].text = col_name

    for _, row in summary_df.iterrows():
        row_cells = t.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val)

    doc.add_heading("3. Desorption Profile Plot", level=1)
    fig_buffer.seek(0)
    doc.add_picture(fig_buffer, width=Inches(5.5))

    doc.add_heading("4. Catalytic Interpretation & Literature Benchmark", level=1)
    for line in conclusions_text:
        doc.add_paragraph(line)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def generate_pdf_report(summary_df, metrics_dict, fig_buffer, conclusions_text):
    """Creates a PDF report via ReportLab."""
    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("<b>CO₂-TPD Characterization Report</b>", styles['Title']))
    story.append(Spacer(1, 12))

    story.append(Paragraph("<b>Key Experimental Parameters:</b>", styles['Heading2']))
    for k, v in metrics_dict.items():
        story.append(Paragraph(f"• <b>{k}:</b> {v}", styles['Normal']))
    story.append(Spacer(1, 12))

    story.append(Paragraph("<b>Basic Site Distribution Breakdown:</b>", styles['Heading2']))
    table_data = [summary_df.columns.tolist()] + summary_df.values.tolist()
    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F77B4')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ]))
    story.append(t)
    story.append(Spacer(1, 12))

    story.append(Paragraph("<b>TPD Profile:</b>", styles['Heading2']))
    fig_buffer.seek(0)
    story.append(RLImage(fig_buffer, width=420, height=280))
    story.append(Spacer(1, 12))

    story.append(Paragraph("<b>Literature Assessment:</b>", styles['Heading2']))
    for line in conclusions_text:
        story.append(Paragraph(line, styles['Normal']))
        story.append(Spacer(1, 4))

    doc.build(story)
    return pdf_buffer.getvalue()


# --- SIDEBAR: 1. DATA INPUT & MAPPING ---
st.sidebar.header("1. Data Input & Mapping")
uploaded_file = st.sidebar.file_uploader("Upload TPD Excel (.xlsx / .xls)", type=["xlsx", "xls"])

if uploaded_file is not None:
    try:
        filename = uploaded_file.name.lower()
        engine = "xlrd" if filename.endswith(".xls") else "openpyxl"

        excel_file = pd.ExcelFile(uploaded_file, engine=engine)
        sheet_name = st.sidebar.selectbox("Select Excel Sheet", excel_file.sheet_names)
        
        raw_full = pd.read_excel(uploaded_file, sheet_name=sheet_name, header=None, engine=engine)
        
        default_header_idx = 23
        for r_idx, row_vals in raw_full.iterrows():
            row_str = " ".join([str(v) for v in row_vals.values if pd.notna(v)])
            if "Temperature" in row_str and "TCD" in row_str:
                default_header_idx = r_idx + 3
                break

        header_row = st.sidebar.number_input("Header Row (0-indexed)", min_value=0, max_value=100, value=default_header_idx)
        df_raw = pd.read_excel(uploaded_file, sheet_name=sheet_name, header=None, engine=engine)
        
        header_vals = df_raw.iloc[header_row].values
        col_options = []
        for i, val in enumerate(header_vals):
            col_name = ""
            temp_i = i
            while temp_i >= 0:
                col_name = chr(temp_i % 26 + 65) + col_name
                temp_i = temp_i // 26 - 1
            
            header_str = str(val) if pd.notna(val) else "Unnamed"
            col_options.append(f"Col {col_name} (Idx {i}): {header_str}")

        def_temp_col = 12 if len(col_options) > 12 else 0
        def_tcd_col = 13 if len(col_options) > 13 else min(1, len(col_options)-1)

        temp_col_sel = st.sidebar.selectbox("Temperature Column", col_options, index=def_temp_col)
        tcd_col_sel = st.sidebar.selectbox("TCD Signal Column", col_options, index=def_tcd_col)
        
        temp_idx = int(temp_col_sel.split("Idx ")[1].split(")")[0])
        tcd_idx = int(tcd_col_sel.split("Idx ")[1].split(")")[0])

        data_df = df_raw.iloc[header_row + 1:].copy()

        sample_mass = st.sidebar.number_input("Sample Mass (mg)", min_value=0.1, value=50.0, step=0.1)
        baseline_subtraction = st.sidebar.checkbox("Apply Linear Baseline Correction", value=True)
        
        # --- SIDEBAR: 2. BASIC SITE BOUNDARIES ---
        st.sidebar.header("2. Basic Site Temperature Boundaries")
        weak_max = st.sidebar.number_input("Weak / Medium Boundary (°C)", min_value=50.0, max_value=300.0, value=200.0, step=10.0)
        medium_max = st.sidebar.number_input("Medium / Strong Boundary (°C)", min_value=200.0, max_value=600.0, value=400.0, step=10.0)

        # --- DATA PRE-PROCESSING ---
        clean_df = pd.DataFrame({
            "Temperature": clean_numeric_series(data_df.iloc[:, temp_idx]),
            "TCD": clean_numeric_series(data_df.iloc[:, tcd_idx])
        }).dropna().sort_values("Temperature").reset_index(drop=True)
        
        if len(clean_df) < 2:
            st.error("Not enough numeric data points found. Adjust 'Header Row' or check column mapping.")
            st.stop()

        clean_df["TCD_Norm"] = (clean_df["TCD"] / sample_mass) * 1000
        
        if baseline_subtraction:
            start_val = clean_df["TCD_Norm"].iloc[0]
            end_val = clean_df["TCD_Norm"].iloc[-1]
            baseline = np.linspace(start_val, end_val, len(clean_df))
            clean_df["TCD_Processed"] = (clean_df["TCD_Norm"] - baseline).clip(lower=0)
        else:
            clean_df["TCD_Processed"] = clean_df["TCD_Norm"]

        # --- REGIONAL INTEGRATION ---
        total_area = simpson(y=clean_df["TCD_Processed"].values, x=clean_df["Temperature"].values)
        
        df_weak = clean_df[clean_df["Temperature"] < weak_max]
        df_med = clean_df[(clean_df["Temperature"] >= weak_max) & (clean_df["Temperature"] < medium_max)]
        df_strong = clean_df[clean_df["Temperature"] >= medium_max]

        area_weak = simpson(y=df_weak["TCD_Processed"].values, x=df_weak["Temperature"].values) if len(df_weak) > 1 else 0.0
        area_med = simpson(y=df_med["TCD_Processed"].values, x=df_med["Temperature"].values) if len(df_med) > 1 else 0.0
        area_strong = simpson(y=df_strong["TCD_Processed"].values, x=df_strong["Temperature"].values) if len(df_strong) > 1 else 0.0

        pct_weak = (area_weak / total_area * 100) if total_area > 0 else 0
        pct_med = (area_med / total_area * 100) if total_area > 0 else 0
        pct_strong = (area_strong / total_area * 100) if total_area > 0 else 0

        # --- SIDEBAR: 3. GRAPHIC & AXIS TYPOGRAPHY CONTROLS ---
        st.sidebar.header("3. Graph Typography & Axes")
        
        with st.sidebar.expander("✏️ Title & Axis Formatting (Bold/Italics/Style)", expanded=True):
            title_text = st.text_input("Plot Title", "CO₂ Temperature-Programmed Desorption Profile")
            title_bold = st.checkbox("Bold Title", value=True)
            title_italic = st.checkbox("Italic Title", value=False)
            title_size = st.slider("Title Size", 8, 24, 14)
            
            xlabel_text = st.text_input("X-Axis Label", "Temperature (°C)")
            xlabel_bold = st.checkbox("Bold X-Label", value=False)
            xlabel_italic = st.checkbox("Italic X-Label", value=False)
            
            ylabel_text = st.text_input("Y-Axis Label", "TCD Signal (a.u. / g_cat)")
            ylabel_bold = st.checkbox("Bold Y-Label", value=False)
            ylabel_italic = st.checkbox("Italic Y-Label", value=False)
            
            label_size = st.slider("Axis Label Size", 8, 20, 12)
            tick_size = st.slider("Tick Label Size", 6, 16, 10)
            font_family = st.selectbox("Font Style", ["DejaVu Sans", "DejaVu Serif", "Arial", "Times New Roman", "Courier New"])

        with st.sidebar.expander("🎨 Color Customization", expanded=False):
            color_mode = st.radio("Shading Style", ["Shade by Basic Site Regions", "Single Color Peak Fill"])
            curve_color = st.color_picker("Main Curve Line Color", "#1F77B4")
            weak_color = st.color_picker("Weak Sites Fill Color", "#3182BD")
            med_color = st.color_picker("Medium Sites Fill Color", "#E6550D")
            strong_color = st.color_picker("Strong Sites Fill Color", "#DE2D26")
            single_fill_color = st.color_picker("Single Fill Color", "#6BAED6")
            fill_alpha = st.slider("Fill Transparency (Alpha)", 0.0, 1.0, 0.35, step=0.05)

        with st.sidebar.expander("📐 Scale Ranges & Rescaling", expanded=False):
            auto_x = st.checkbox("Auto X-Axis Range", value=True)
            min_temp, max_temp = float(clean_df["Temperature"].min()), float(clean_df["Temperature"].max())
            if not auto_x:
                x_min = st.number_input("X Min (°C)", value=min_temp, step=10.0)
                x_max = st.number_input("X Max (°C)", value=max_temp, step=10.0)
            
            auto_y = st.checkbox("Auto Y-Axis Range", value=True)
            min_tcd, max_tcd = float(clean_df["TCD_Processed"].min()), float(clean_df["TCD_Processed"].max())
            if not auto_y:
                y_min = st.number_input("Y Min", value=0.0, step=0.01)
                y_max = st.number_input("Y Max", value=max_tcd * 1.1, step=0.01)

        with st.sidebar.expander("📏 Ticks, Lines & Canvas Options", expanded=False):
            show_major_ticks = st.checkbox("Show Major Ticks", value=True)
            show_minor_ticks = st.checkbox("Show Minor Ticks", value=True)
            show_grid = st.checkbox("Show Grid Lines", value=False)
            show_box_border = st.checkbox("Box Frame (Top/Right Spines)", value=True)
            line_width = st.slider("Line Width (pt)", 0.5, 4.0, 1.5, step=0.25)
            fig_width = st.slider("Figure Width (in)", 3.0, 10.0, 6.5, step=0.5)
            fig_height = st.slider("Figure Height (in)", 2.5, 8.0, 4.5, step=0.5)
            dpi_val = st.number_input("Export DPI", value=600, step=100)

        # Helper to convert style toggles to Matplotlib font weight/style
        def get_font_style(is_bold, is_italic):
            weight = "bold" if is_bold else "normal"
            style = "italic" if is_italic else "normal"
            return weight, style

        # --- PLOTTING ENGINE ---
        plt.rcParams["font.family"] = font_family
        fig, ax = plt.subplots(figsize=(fig_width, fig_height), dpi=150)

        ax.plot(clean_df["Temperature"], clean_df["TCD_Processed"], color=curve_color, linewidth=line_width, label="CO₂ Signal")

        if color_mode == "Shade by Basic Site Regions":
            ax.fill_between(df_weak["Temperature"], df_weak["TCD_Processed"], color=weak_color, alpha=fill_alpha, label=f"Weak (<{weak_max:.0f}°C)")
            ax.fill_between(df_med["Temperature"], df_med["TCD_Processed"], color=med_color, alpha=fill_alpha, label=f"Medium ({weak_max:.0f}-{medium_max:.0f}°C)")
            ax.fill_between(df_strong["Temperature"], df_strong["TCD_Processed"], color=strong_color, alpha=fill_alpha, label=f"Strong (>{medium_max:.0f}°C)")
            ax.legend(fontsize=tick_size - 1, frameon=False)
        else:
            ax.fill_between(clean_df["Temperature"], clean_df["TCD_Processed"], color=single_fill_color, alpha=fill_alpha)

        # Apply customized font weights/styles to titles & labels
        t_weight, t_style = get_font_style(title_bold, title_italic)
        x_weight, x_style = get_font_style(xlabel_bold, xlabel_italic)
        y_weight, y_style = get_font_style(ylabel_bold, ylabel_italic)

        ax.set_title(title_text, fontsize=title_size, fontweight=t_weight, fontstyle=t_style)
        ax.set_xlabel(xlabel_text, fontsize=label_size, fontweight=x_weight, fontstyle=x_style)
        ax.set_ylabel(ylabel_text, fontsize=label_size, fontweight=y_weight, fontstyle=y_style)

        # Axis scaling
        if not auto_x:
            ax.set_xlim(x_min, x_max)
        if not auto_y:
            ax.set_ylim(y_min, y_max)

        # Ticks and Grid toggles
        if show_major_ticks:
            ax.tick_params(axis="both", which="major", labelsize=tick_size, bottom=True, left=True)
        else:
            ax.tick_params(axis="both", which="major", bottom=False, left=False, labelbottom=False, labelleft=False)

        if show_minor_ticks and show_major_ticks:
            ax.xaxis.set_minor_locator(AutoMinorLocator())
            ax.yaxis.set_minor_locator(AutoMinorLocator())
            ax.tick_params(axis="both", which="minor", bottom=True, left=True)
        else:
            ax.tick_params(axis="both", which="minor", bottom=False, left=False)

        if show_grid:
            ax.grid(True, linestyle="--", alpha=0.5)

        if not show_box_border:
            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)

        # --- MAIN DISPLAY LAYOUT ---
        col1, col2 = st.columns([1.6, 1])

        with col1:
            st.pyplot(fig)
            img_buffer = io.BytesIO()
            fig.savefig(img_buffer, format="png", dpi=dpi_val, bbox_inches="tight")
            st.download_button(
                label="📥 Download High-Res Plot (PNG)",
                data=img_buffer.getvalue(),
                file_name="CO2_TPD_Plot.png",
                mime="image/png"
            )

        with col2:
            st.subheader("📊 Quantified Basic Sites")
            summary_table = pd.DataFrame({
                "Site Type": ["Weak Sites", "Medium Sites", "Strong Sites", "Total Basicity"],
                "Temp Range": [f"< {weak_max:.0f} °C", f"{weak_max:.0f} – {medium_max:.0f} °C", f"> {medium_max:.0f} °C", "Full Spectrum"],
                "Desorption Area": [f"{area_weak:.2f}", f"{area_med:.2f}", f"{area_strong:.2f}", f"{total_area:.2f}"],
                "Fraction (%)": [f"{pct_weak:.1f} %", f"{pct_med:.1f} %", f"{pct_strong:.1f} %", "100.0 %"]
            })
            st.dataframe(summary_table, hide_index=True, use_container_width=True)

            max_idx = clean_df["TCD_Processed"].idxmax()
            peak_temp = clean_df.loc[max_idx, "Temperature"]
            st.metric("Primary Peak Temp (T_max)", f"{peak_temp:.1f} °C")
            st.metric("Sample Mass Evaluated", f"{sample_mass} mg")

        # --- LITERATURE COMPARISON MODULE ---
        st.markdown("---")
        st.header("🧠 Literature Comparison & Methanol Synthesis Correlation")

        cat_col1, cat_col2 = st.columns(2)
        conclusions_list = []

        with cat_col1:
            st.subheader("📌 Main Conclusions from Profile")
            dominant_site = "Weak" if pct_weak >= max(pct_med, pct_strong) else ("Medium" if pct_med >= pct_strong else "Strong")
            
            c1 = f"The catalyst profile is dominated by {dominant_site} Basic Sites ({max(pct_weak, pct_med, pct_strong):.1f}% of total surface basicity)."
            c2 = f"Primary desorption peak maximum (T_max) occurs at {peak_temp:.1f} °C."
            st.write(f"• **Dominant Surface Species:** {c1}")
            st.write(f"• **Desorption Thermal Maxima:** {c2}")
            conclusions_list.extend([c1, c2])

        with cat_col2:
            st.subheader("💡 Correlation to CO₂/CO/H₂ Methanol Synthesis")
            if pct_med >= 40.0:
                c3 = "HIGH CATALYTIC POTENTIAL: Medium basic sites (200-400 °C) are active sites for optimal CO2 activation, stabilizing bidentate formate (*HCOO) intermediates for hydrogenation to methanol."
                st.success(f"🟢 **{c3}**")
            elif pct_weak > 50.0:
                c3 = "MODERATE ACTIVITY: High weak sites (<200 °C) concentration leads to quick CO2 desorption prior to reaction."
                st.warning(f"🟡 **{c3}**")
            else:
                c3 = "CARBONATE POISONING RISK: Excessive strong basic sites (>400 °C) hold CO2 rigidly as monodentate carbonates, promoting side reactions like RWGS to CO."
                st.error(f"🔴 **{c3}**")
            conclusions_list.append(c3)

        # --- EXPORT REPORT SECTION ---
        st.markdown("---")
        st.header("📄 Export Characterization Report")
        
        metrics_summary = {
            "Sample Mass": f"{sample_mass} mg",
            "Primary Peak (T_max)": f"{peak_temp:.1f} °C",
            "Total Basicity Area": f"{total_area:.2f} a.u.*°C/g",
            "Dominant Site Class": dominant_site
        }

        exp_col1, exp_col2, exp_col3 = st.columns(3)

        with exp_col1:
            pdf_bytes = generate_pdf_report(summary_table, metrics_summary, img_buffer, conclusions_list)
            st.download_button("📥 Export Report as PDF", pdf_bytes, "CO2_TPD_Report.pdf", "application/pdf")

        with exp_col2:
            word_bytes = generate_word_report(summary_table, metrics_summary, img_buffer, conclusions_list)
            st.download_button("📥 Export Report as Word (.docx)", word_bytes, "CO2_TPD_Report.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

        with exp_col3:
            excel_bytes = generate_excel_report(summary_table, metrics_summary, img_buffer)
            st.download_button("📥 Export Data & Plot as Excel (.xlsx)", excel_bytes, "CO2_TPD_Report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    except Exception as e:
        st.error(f"Error executing analysis: {e}")
else:
    st.info("👋 Upload a CO₂-TPD dataset via the sidebar to access custom plotting, typography formatting, and report export tools.")
